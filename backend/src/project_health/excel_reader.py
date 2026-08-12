from __future__ import annotations

import re
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Any

NS = {"a": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
RELS_NS = {
    "r": "http://schemas.openxmlformats.org/package/2006/relationships",
    "od": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}


def read_workbook(path: str | Path) -> dict[str, list[dict[str, Any]]]:
    """Read an xlsx workbook into sheet-name keyed row dictionaries.

    Uses openpyxl when available. Falls back to a small OOXML reader so the
    backend can still run in restricted review environments.
    """
    try:
        return _read_with_openpyxl(Path(path))
    except ModuleNotFoundError:
        return _read_with_ooxml(Path(path))


def _read_with_openpyxl(path: Path) -> dict[str, list[dict[str, Any]]]:
    from openpyxl import load_workbook

    sheets: dict[str, list[dict[str, Any]]] = {}
    wb = load_workbook(path, data_only=True)
    try:
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                sheets[ws.title] = []
                continue
            if _raw_row_sheet(ws.title):
                records = []
                for offset, values in enumerate(rows, start=1):
                    if not any(v is not None and str(v).strip() for v in values):
                        continue
                    row = {"__row_number__": offset}
                    for idx, value in enumerate(values, start=1):
                        row[f"col_{idx}"] = value
                    records.append(row)
                sheets[ws.title] = records
                continue
            headers = [_clean_header(v) for v in rows[0]]
            records = []
            for offset, values in enumerate(rows[1:], start=2):
                if not any(v is not None and str(v).strip() for v in values):
                    continue
                row = {"__row_number__": offset}
                for header, value in zip(headers, values):
                    if header:
                        row[header] = value
                records.append(row)
            sheets[ws.title] = records
    finally:
        try:
            wb.close()
        except AttributeError:
            pass
    return sheets


def _read_with_ooxml(path: Path) -> dict[str, list[dict[str, Any]]]:
    with zipfile.ZipFile(path) as zf:
        shared = _shared_strings(zf)
        sheet_paths = _sheet_paths(zf)
        result: dict[str, list[dict[str, Any]]] = {}
        for sheet_name, sheet_path in sheet_paths:
            root = ET.fromstring(zf.read(sheet_path))
            rows = []
            for row in root.findall(".//a:sheetData/a:row", NS):
                values_by_col: dict[int, Any] = {}
                for cell in row.findall("a:c", NS):
                    values_by_col[_column_index(cell.get("r", ""))] = _cell_value(cell, shared)
                rows.append((int(row.get("r", len(rows) + 1)), values_by_col))
            if not rows:
                result[sheet_name] = []
                continue
            if _raw_row_sheet(sheet_name):
                records = []
                for row_number, values_by_col in rows:
                    if not any(_present(v) for v in values_by_col.values()):
                        continue
                    record = {"__row_number__": row_number}
                    for idx, value in values_by_col.items():
                        record[f"col_{idx}"] = value
                    records.append(record)
                result[sheet_name] = records
                continue
            headers = {idx: _clean_header(value) for idx, value in rows[0][1].items()}
            records = []
            for row_number, values_by_col in rows[1:]:
                if not any(_present(v) for v in values_by_col.values()):
                    continue
                record = {"__row_number__": row_number}
                for idx, value in values_by_col.items():
                    header = headers.get(idx)
                    if header:
                        record[header] = value
                records.append(record)
            result[sheet_name] = records
        return result


def _shared_strings(zf: zipfile.ZipFile) -> list[str]:
    if "xl/sharedStrings.xml" not in zf.namelist():
        return []
    root = ET.fromstring(zf.read("xl/sharedStrings.xml"))
    return ["".join(si.itertext()).strip() for si in root.findall("a:si", NS)]


def _sheet_paths(zf: zipfile.ZipFile) -> list[tuple[str, str]]:
    workbook = ET.fromstring(zf.read("xl/workbook.xml"))
    rels = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
    rel_map = {rel.get("Id"): rel.get("Target") for rel in rels.findall("r:Relationship", RELS_NS)}
    sheets = []
    for sheet in workbook.findall("a:sheets/a:sheet", NS):
        rid = sheet.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
        target = rel_map.get(rid, "")
        if not target.startswith("/"):
            target = "xl/" + target
        else:
            target = target.lstrip("/")
        sheets.append((sheet.get("name", target), target))
    return sheets


def _cell_value(cell: ET.Element, shared: list[str]) -> Any:
    cell_type = cell.get("t")
    value = cell.find("a:v", NS)
    if cell_type == "s" and value is not None:
        return shared[int(value.text or 0)]
    if cell_type == "inlineStr":
        return "".join(cell.itertext()).strip()
    if value is None:
        return None
    text = value.text
    if text is None:
        return None
    if re.fullmatch(r"-?\d+", text):
        return int(text)
    if re.fullmatch(r"-?\d+\.\d+", text):
        return float(text)
    return text


def _column_index(cell_ref: str) -> int:
    letters = re.match(r"[A-Z]+", cell_ref)
    if not letters:
        return 0
    value = 0
    for char in letters.group(0):
        value = value * 26 + (ord(char) - ord("A") + 1)
    return value


def _clean_header(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def _raw_row_sheet(sheet_name: str) -> bool:
    lowered = sheet_name.lower().strip()
    return lowered == "summary" or "comment" in lowered


def _present(value: Any) -> bool:
    return value is not None and str(value).strip() != ""
