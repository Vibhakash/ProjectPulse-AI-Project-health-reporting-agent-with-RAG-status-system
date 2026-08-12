from __future__ import annotations

import json
import sqlite3
from collections import Counter
from datetime import date
from pathlib import Path
from typing import Any

from .repositories import fetch_comment_themes, fetch_portfolio, fetch_snapshots, fetch_top_risk_tasks


def generate_monthly_deck(conn: sqlite3.Connection, output_dir: str | Path) -> str:
    output_path = Path(output_dir)
    output_path.mkdir(parents=True, exist_ok=True)
    pptx_path = output_path / "monthly_project_health.pptx"

    try:
        _generate_with_python_pptx(conn, pptx_path)
    except ModuleNotFoundError:
        fallback = output_path / "monthly_project_health_deck.md"
        fallback.write_text(_deck_markdown(conn), encoding="utf-8")
        return str(fallback)
    return str(pptx_path)


def _load_json_data(weekly_json_path: str | None) -> dict[str, Any]:
    if not weekly_json_path:
        return {}
    p = Path(weekly_json_path)
    if not p.exists():
        # Try checking relative to outputs directory or package
        p = Path("outputs") / "weekly" / p.name
        if not p.exists():
            return {}
    try:
        with open(p, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return {}


def _generate_with_python_pptx(conn: sqlite3.Connection, pptx_path: Path) -> None:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    portfolio = fetch_portfolio(conn)
    snapshots = fetch_snapshots(conn)
    risks = fetch_top_risk_tasks(conn, 8)
    comments = fetch_comment_themes(conn, 20)
    rag_counts = Counter(row["rag_status"] for row in portfolio)

    # Cohesive, premium design color system
    c_dark_slate = RGBColor(15, 23, 42)
    c_medium_slate = RGBColor(71, 85, 105)
    c_light_bg = RGBColor(248, 250, 252)
    c_white = RGBColor(255, 255, 255)
    c_border = RGBColor(226, 232, 240)
    c_accent_indigo = RGBColor(99, 102, 241)

    c_green = RGBColor(16, 185, 129)
    c_amber = RGBColor(245, 158, 11)
    c_red = RGBColor(239, 68, 68)

    def get_rag_color(status):
        return {"Red": c_red, "Amber": c_amber, "Green": c_green}.get(status, c_medium_slate)

    def format_tf(tf):
        tf.word_wrap = True
        tf.margin_left = Inches(0.2)
        tf.margin_right = Inches(0.2)
        tf.margin_top = Inches(0.2)
        tf.margin_bottom = Inches(0.2)

    def add_card(slide, left, top, width, height, bg_color=c_white, border_color=c_border):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.5)
        if shape.adjustments:
            try:
                shape.adjustments[0] = 0.05
            except Exception:
                pass
        return shape

    def add_header(slide, title, subtitle=None, category=None):
        if category:
            box_cat = slide.shapes.add_textbox(Inches(0.55), Inches(0.2), Inches(12.2), Inches(0.3))
            tf_cat = box_cat.text_frame
            tf_cat.word_wrap = True
            p_cat = tf_cat.paragraphs[0]
            p_cat.text = category.upper()
            p_cat.font.name = "Arial"
            p_cat.font.size = Pt(9)
            p_cat.font.bold = True
            p_cat.font.color.rgb = c_accent_indigo

        box = slide.shapes.add_textbox(Inches(0.55), Inches(0.4), Inches(12.2), Inches(0.8))
        tf = box.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.name = "Arial"
        p1.font.size = Pt(24)
        p1.font.bold = True
        p1.font.color.rgb = c_dark_slate

        if subtitle:
            p2 = tf.add_paragraph()
            p2.text = subtitle
            p2.font.name = "Arial"
            p2.font.size = Pt(11)
            p2.font.color.rgb = c_medium_slate
            p2.space_before = Pt(3)

    # ----------------------------------------------------
    # Slide 1: Title Slide
    # ----------------------------------------------------
    slide_title = prs.slides.add_slide(prs.slide_layouts[6])
    top_band = slide_title.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.15))
    top_band.fill.solid()
    top_band.fill.fore_color.rgb = c_accent_indigo
    top_band.line.fill.background()

    title_box = slide_title.shapes.add_textbox(Inches(0.75), Inches(1.2), Inches(11.83), Inches(1.8))
    tf = title_box.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = "PROJECT PORTFOLIO HEALTH SYNTHESIS"
    p1.font.name = "Arial"
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = c_dark_slate

    p2 = tf.add_paragraph()
    p2.text = "Monthly Executive Performance Review & Automated RAG Analysis"
    p2.font.name = "Arial"
    p2.font.size = Pt(16)
    p2.font.color.rgb = c_medium_slate
    p2.space_before = Pt(8)

    p3 = tf.add_paragraph()
    p3.text = f"Generated by Project Health AI Agent | Date: {date.today().isoformat()}"
    p3.font.name = "Arial"
    p3.font.size = Pt(11)
    p3.font.color.rgb = c_accent_indigo
    p3.space_before = Pt(16)

    statuses_info = [
        ("Green", "On Track", c_green, "Projects executing smoothly against baseline schedule and budget expectations."),
        ("Amber", "Under Watch", c_amber, "Projects exhibiting minor delays, progress gaps, or warning indicators."),
        ("Red", "Critical Escalate", c_red, "Projects with significant delays, blocked milestones, or high risk items.")
    ]

    for idx, (status, label, color, desc) in enumerate(statuses_info):
        left = Inches(0.75 + idx * 4.0)
        top = Inches(3.6)
        width = Inches(3.7)
        height = Inches(2.7)

        card = add_card(slide_title, left, top, width, height)
        tf_card = card.text_frame
        format_tf(tf_card)

        p_label = tf_card.paragraphs[0]
        p_label.text = f"{label.upper()} / {status.upper()}"
        p_label.font.name = "Arial"
        p_label.font.size = Pt(10)
        p_label.font.bold = True
        p_label.font.color.rgb = color

        p_num = tf_card.add_paragraph()
        p_num.text = str(rag_counts.get(status, 0))
        p_num.font.name = "Arial"
        p_num.font.size = Pt(44)
        p_num.font.bold = True
        p_num.font.color.rgb = c_dark_slate
        p_num.space_before = Pt(4)

        p_desc = tf_card.add_paragraph()
        p_desc.text = desc
        p_desc.font.name = "Arial"
        p_desc.font.size = Pt(10.5)
        p_desc.font.color.rgb = c_medium_slate
        p_desc.space_before = Pt(12)

    # ----------------------------------------------------
    # Slide 2: Portfolio Directory Table
    # ----------------------------------------------------
    slide_dir = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide_dir, "Portfolio Health Directory", "Consolidated view of active project stages, scores, and data confidence", "Portfolio Summary")

    left = Inches(0.55)
    top = Inches(1.4)
    width = Inches(12.23)
    row_count = len(portfolio) + 1
    height = Inches(min(5.2, 0.45 * row_count))

    table_shape = slide_dir.shapes.add_table(row_count, 6, left, top, width, height)
    table = table_shape.table

    table.columns[0].width = Inches(3.0)
    table.columns[1].width = Inches(1.3)
    table.columns[2].width = Inches(1.2)
    table.columns[3].width = Inches(1.4)
    table.columns[4].width = Inches(3.6)
    table.columns[5].width = Inches(1.73)

    headers_text = ["Project Name", "RAG Status", "RAG Score", "Data Quality", "Current Stage", "Project Manager"]
    for col_idx, h_text in enumerate(headers_text):
        cell = table.cell(0, col_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = c_dark_slate
        tf_cell = cell.text_frame
        tf_cell.word_wrap = True
        p = tf_cell.paragraphs[0]
        p.text = h_text
        p.alignment = PP_ALIGN.LEFT if col_idx in (0, 4, 5) else PP_ALIGN.CENTER
        p.font.name = "Arial"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = c_white

    for row_idx, row in enumerate(portfolio, start=1):
        bg = c_light_bg if row_idx % 2 == 1 else c_white

        cell_name = table.cell(row_idx, 0)
        cell_name.fill.solid()
        cell_name.fill.fore_color.rgb = bg
        p_name = cell_name.text_frame.paragraphs[0]
        p_name.text = str(row["name"])
        p_name.font.name = "Arial"
        p_name.font.size = Pt(11)
        p_name.font.bold = True
        p_name.font.color.rgb = c_dark_slate

        cell_status = table.cell(row_idx, 1)
        cell_status.fill.solid()
        cell_status.fill.fore_color.rgb = bg
        p_status = cell_status.text_frame.paragraphs[0]
        p_status.text = str(row["rag_status"])
        p_status.alignment = PP_ALIGN.CENTER
        p_status.font.name = "Arial"
        p_status.font.size = Pt(11)
        p_status.font.bold = True
        p_status.font.color.rgb = get_rag_color(row["rag_status"])

        cell_score = table.cell(row_idx, 2)
        cell_score.fill.solid()
        cell_score.fill.fore_color.rgb = bg
        p_score = cell_score.text_frame.paragraphs[0]
        p_score.text = f"{row['rag_score']}/100"
        p_score.alignment = PP_ALIGN.CENTER
        p_score.font.name = "Arial"
        p_score.font.size = Pt(11)
        p_score.font.color.rgb = c_medium_slate

        cell_dq = table.cell(row_idx, 3)
        cell_dq.fill.solid()
        cell_dq.fill.fore_color.rgb = bg
        p_dq = cell_dq.text_frame.paragraphs[0]
        p_dq.text = f"{int(row['data_quality_score'])}%"
        p_dq.alignment = PP_ALIGN.CENTER
        p_dq.font.name = "Arial"
        p_dq.font.size = Pt(11)
        p_dq.font.color.rgb = c_medium_slate

        cell_stage = table.cell(row_idx, 4)
        cell_stage.fill.solid()
        cell_stage.fill.fore_color.rgb = bg
        p_stage = cell_stage.text_frame.paragraphs[0]
        p_stage.text = str(row["project_stage"] or "Unknown")
        p_stage.font.name = "Arial"
        p_stage.font.size = Pt(11)
        p_stage.font.color.rgb = c_medium_slate

        cell_pm = table.cell(row_idx, 5)
        cell_pm.fill.solid()
        cell_pm.fill.fore_color.rgb = bg
        p_pm = cell_pm.text_frame.paragraphs[0]
        p_pm.text = str(row["project_manager"] or "Unknown")
        p_pm.font.name = "Arial"
        p_pm.font.size = Pt(11)
        p_pm.font.color.rgb = c_medium_slate

    # ----------------------------------------------------
    # Slide 3: Portfolio Narrative & Sentiment (Consolidated)
    # ----------------------------------------------------
    slide_narrative = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide_narrative, "Portfolio Narrative & Sentiment Summary", "Consolidated AI summaries and stakeholder sentiment comments across all projects", "Portfolio Summary")

    # Left Column Card: AI Narrative
    card_n1 = add_card(slide_narrative, Inches(0.55), Inches(1.35), Inches(5.9), Inches(5.2))
    tf_n1 = card_n1.text_frame
    format_tf(tf_n1)

    p_n1_hdr = tf_n1.paragraphs[0]
    p_n1_hdr.text = "PORTFOLIO EXECUTIVE NARRATIVES"
    p_n1_hdr.font.name = "Arial"
    p_n1_hdr.font.size = Pt(12)
    p_n1_hdr.font.bold = True
    p_n1_hdr.font.color.rgb = c_accent_indigo
    p_n1_hdr.space_after = Pt(10)

    # Order projects: Red/Amber first
    red_amber_projects = [r for r in portfolio if r["rag_status"] in ("Red", "Amber")]
    green_projects = [r for r in portfolio if r["rag_status"] == "Green"]
    ordered_projects = red_amber_projects + green_projects

    if ordered_projects:
        # Show top projects narratives, limit to fit space
        for row in ordered_projects[:4]:
            p_name = row["name"]
            status = row["rag_status"]
            json_data = _load_json_data(row["weekly_json_path"])
            rag_data = json_data.get("rag", {})
            exec_sum = rag_data.get("executive_summary") or ""
            if not exec_sum:
                reasons = rag_data.get("reasons") or []
                exec_sum = reasons[0] if reasons else "No executive summary provided."
            
            p_proj = tf_n1.add_paragraph()
            p_proj.space_after = Pt(6)
            run_name = p_proj.add_run()
            run_name.text = f"• {p_name} ({status}): "
            run_name.font.name = "Arial"
            run_name.font.bold = True
            run_name.font.size = Pt(10)
            run_name.font.color.rgb = get_rag_color(status)

            run_text = p_proj.add_run()
            max_len = 160
            if len(exec_sum) > max_len:
                exec_sum = exec_sum[:max_len] + "..."
            run_text.text = exec_sum
            run_text.font.name = "Arial"
            run_text.font.size = Pt(10)
            run_text.font.color.rgb = c_dark_slate
    else:
        p_empty = tf_n1.add_paragraph()
        p_empty.text = "• No project data available."
        p_empty.font.name = "Arial"
        p_empty.font.size = Pt(10.5)
        p_empty.font.color.rgb = c_medium_slate

    # Right Column Card: Sentiment & Themes
    card_n2 = add_card(slide_narrative, Inches(6.88), Inches(1.35), Inches(5.9), Inches(5.2))
    tf_n2 = card_n2.text_frame
    format_tf(tf_n2)

    p_n2_hdr = tf_n2.paragraphs[0]
    p_n2_hdr.text = "STAKEHOLDER SENTIMENT & THEMES"
    p_n2_hdr.font.name = "Arial"
    p_n2_hdr.font.size = Pt(12)
    p_n2_hdr.font.bold = True
    p_n2_hdr.font.color.rgb = c_accent_indigo
    p_n2_hdr.space_after = Pt(10)

    # Pull sentiment summaries
    sentiment_count = 0
    for row in ordered_projects[:3]:
        p_name = row["name"]
        status = row["rag_status"]
        json_data = _load_json_data(row["weekly_json_path"])
        rag_data = json_data.get("rag", {})
        sent_sum = rag_data.get("sentiment_summary") or ""
        r_themes = rag_data.get("risk_themes") or []

        if sent_sum:
            p_sent = tf_n2.add_paragraph()
            p_sent.space_after = Pt(4)
            run_name = p_sent.add_run()
            run_name.text = f"• {p_name}: "
            run_name.font.name = "Arial"
            run_name.font.bold = True
            run_name.font.size = Pt(10)
            run_name.font.color.rgb = get_rag_color(status)

            run_text = p_sent.add_run()
            max_len = 120
            if len(sent_sum) > max_len:
                sent_sum = sent_sum[:max_len] + "..."
            run_text.text = sent_sum
            run_text.font.name = "Arial"
            run_text.font.size = Pt(10)
            run_text.font.color.rgb = c_medium_slate
            sentiment_count += 1

            if r_themes:
                p_themes = tf_n2.add_paragraph()
                p_themes.space_after = Pt(6)
                run_t = p_themes.add_run()
                run_t.text = f"  Key Risk Themes: {', '.join(r_themes[:3])}"
                run_t.font.name = "Arial"
                run_t.font.italic = True
                run_t.font.size = Pt(9.5)
                run_t.font.color.rgb = c_accent_indigo

    if sentiment_count == 0:
        p_empty = tf_n2.add_paragraph()
        p_empty.text = "• No feedback or comments detected across the portfolio projects."
        p_empty.font.name = "Arial"
        p_empty.font.size = Pt(10.5)
        p_empty.font.color.rgb = c_medium_slate

    # ----------------------------------------------------
    # Slide 4: Systemic Risks & Signals (Consolidated)
    # ----------------------------------------------------
    slide_risks = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide_risks, "Emerging Systemic Risks & Signal Analysis", "Critical task slippages and cross-project signal alert distributions", "Portfolio Summary")

    # Left Column Card: Critical Task slippages
    card_r1 = add_card(slide_risks, Inches(0.55), Inches(1.35), Inches(5.9), Inches(5.2))
    tf_r1 = card_r1.text_frame
    format_tf(tf_r1)

    p_r1_hdr = tf_r1.paragraphs[0]
    p_r1_hdr.text = "TOP DELIVERY RISKS & CRITICAL TASK EVIDENCE"
    p_r1_hdr.font.name = "Arial"
    p_r1_hdr.font.size = Pt(12)
    p_r1_hdr.font.bold = True
    p_r1_hdr.font.color.rgb = c_accent_indigo
    p_r1_hdr.space_after = Pt(10)

    # Fetch top tasks at risk across all projects
    portfolio_risks = fetch_top_risk_tasks(conn, 5)
    if portfolio_risks:
        for ar in portfolio_risks:
            p_ar = tf_r1.add_paragraph()
            p_ar.space_after = Pt(8)
            
            run_p = p_ar.add_run()
            run_p.text = f"• {ar['name']} (Row {ar['source_row']}): "
            run_p.font.name = "Arial"
            run_p.font.bold = True
            run_p.font.size = Pt(9.5)
            run_p.font.color.rgb = get_rag_color(ar['rag_status'])

            run_t = p_ar.add_run()
            run_t.text = f"{ar['task_name']}"
            run_t.font.name = "Arial"
            run_t.font.bold = True
            run_t.font.size = Pt(9.5)
            run_t.font.color.rgb = c_dark_slate

            p_details = tf_r1.add_paragraph()
            p_details.space_after = Pt(4)
            p_details.font.name = "Arial"
            p_details.font.size = Pt(9)
            p_details.font.color.rgb = c_medium_slate
            
            status_str = ar['status'] or "N/A"
            float_str = f"Float: {ar['total_float']}d" if ar['total_float'] is not None else "Float: N/A"
            health_str = f"Health: {ar['source_schedule_health']}" if ar['source_schedule_health'] else "Health: N/A"
            p_details.text = f"  Status: {status_str} | {health_str} | {float_str} | Comp: {ar['percent_complete'] or 0}%"
            if ar['status_comment']:
                p_comm = tf_r1.add_paragraph()
                p_comm.space_after = Pt(6)
                run_comm = p_comm.add_run()
                run_comm.text = f"  Comment: {ar['status_comment']}"
                run_comm.font.name = "Arial"
                run_comm.font.italic = True
                run_comm.font.size = Pt(8.5)
                run_comm.font.color.rgb = c_medium_slate
    else:
        p_empty = tf_r1.add_paragraph()
        p_empty.text = "• No critical path or delayed task evidence found across the portfolio."
        p_empty.font.name = "Arial"
        p_empty.font.size = Pt(10.5)
        p_empty.font.color.rgb = c_medium_slate

    # Right Column Card: Elevated Signals & Alerts
    card_r2 = add_card(slide_risks, Inches(6.88), Inches(1.35), Inches(5.9), Inches(5.2))
    tf_r2 = card_r2.text_frame
    format_tf(tf_r2)

    p_r2_hdr = tf_r2.paragraphs[0]
    p_r2_hdr.text = "ELEVATED SIGNAL & FLIP ALERTS"
    p_r2_hdr.font.name = "Arial"
    p_r2_hdr.font.size = Pt(12)
    p_r2_hdr.font.bold = True
    p_r2_hdr.font.color.rgb = c_accent_indigo
    p_r2_hdr.space_after = Pt(10)

    # RAG flip alerts
    flips = [row["rag_flip_alert"] for row in snapshots if row["rag_flip_alert"]]
    if flips:
        p_flip_hdr = tf_r2.add_paragraph()
        p_flip_hdr.text = "RAG STATUS CHANGE ALERTS:"
        p_flip_hdr.font.name = "Arial"
        p_flip_hdr.font.size = Pt(10)
        p_flip_hdr.font.bold = True
        p_flip_hdr.font.color.rgb = c_red
        p_flip_hdr.space_after = Pt(4)

        for flip in flips[:3]:
            p_flip = tf_r2.add_paragraph()
            p_flip.space_after = Pt(6)
            run_f = p_flip.add_run()
            run_f.text = f"• {flip}"
            run_f.font.name = "Arial"
            run_f.font.size = Pt(9.5)
            run_f.font.bold = True
            run_f.font.color.rgb = c_red
    else:
        p_no_flip = tf_r2.add_paragraph()
        p_no_flip.text = "• No recent RAG status flips detected."
        p_no_flip.font.name = "Arial"
        p_no_flip.font.size = Pt(9.5)
        p_no_flip.font.color.rgb = c_medium_slate
        p_no_flip.space_after = Pt(6)

    # Calculate elevated signals
    latest_snapshot_ids = [row["id"] for row in portfolio]
    elevated_counts = Counter()
    if latest_snapshot_ids:
        placeholders = ",".join("?" * len(latest_snapshot_ids))
        active_signals = conn.execute(
            f"SELECT signal_name, score, triggered_override FROM rag_signals WHERE snapshot_id IN ({placeholders})",
            latest_snapshot_ids
        ).fetchall()
        for rs in active_signals:
            if rs["score"] >= 60 or rs["triggered_override"]:
                elevated_counts[rs["signal_name"]] += 1

    tf_r2.add_paragraph().space_after = Pt(10)
    p_sig_hdr = tf_r2.add_paragraph()
    p_sig_hdr.text = "PORTFOLIO SIGNAL ALERT FREQUENCY:"
    p_sig_hdr.font.name = "Arial"
    p_sig_hdr.font.size = Pt(10)
    p_sig_hdr.font.bold = True
    p_sig_hdr.font.color.rgb = c_dark_slate
    p_sig_hdr.space_after = Pt(6)

    if elevated_counts:
        for sig_name, count in elevated_counts.items():
            p_sig = tf_r2.add_paragraph()
            p_sig.space_after = Pt(4)
            run_s = p_sig.add_run()
            run_s.text = f"• {sig_name}: "
            run_s.font.name = "Arial"
            run_s.font.bold = True
            run_s.font.size = Pt(9.5)
            run_s.font.color.rgb = c_dark_slate

            run_c = p_sig.add_run()
            run_c.text = f"Elevated in {count} project(s)"
            run_c.font.name = "Arial"
            run_c.font.size = Pt(9.5)
            run_c.font.color.rgb = c_medium_slate
    else:
        p_no_sig = tf_r2.add_paragraph()
        p_no_sig.text = "• No elevated signal thresholds detected across active projects."
        p_no_sig.font.name = "Arial"
        p_no_sig.font.size = Pt(9.5)
        p_no_sig.font.color.rgb = c_medium_slate

    # ----------------------------------------------------
    # Slide 5: Recommended Leadership Actions (Consolidated)
    # ----------------------------------------------------
    slide_actions = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide_actions, "Recommended Leadership & Escalation Actions", "Immediate interventions for at-risk projects and general portfolio quality recommendations", "Portfolio Summary")

    # Left Column Card: Immediate Escalations
    card_e1 = add_card(slide_actions, Inches(0.55), Inches(1.35), Inches(5.9), Inches(5.2))
    tf_e1 = card_e1.text_frame
    format_tf(tf_e1)

    p_e1_hdr = tf_e1.paragraphs[0]
    p_e1_hdr.text = "IMMEDIATE PROJECT ESCALATIONS"
    p_e1_hdr.font.name = "Arial"
    p_e1_hdr.font.size = Pt(12)
    p_e1_hdr.font.bold = True
    p_e1_hdr.font.color.rgb = c_accent_indigo
    p_e1_hdr.space_after = Pt(10)

    escalations = []
    for row in ordered_projects:
        if row["rag_status"] in ("Red", "Amber"):
            json_data = _load_json_data(row["weekly_json_path"])
            recs = json_data.get("rag", {}).get("recommendations") or []
            for rec in recs:
                escalations.append((row["name"], row["rag_status"], rec))

    if escalations:
        for proj_name, status, rec in escalations[:4]:
            p_esc = tf_e1.add_paragraph()
            p_esc.space_after = Pt(8)
            run_p = p_esc.add_run()
            run_p.text = f"• {proj_name} ({status}): "
            run_p.font.name = "Arial"
            run_p.font.bold = True
            run_p.font.size = Pt(10)
            run_p.font.color.rgb = get_rag_color(status)

            run_r = p_esc.add_run()
            run_r.text = rec
            run_r.font.name = "Arial"
            run_r.font.size = Pt(10)
            run_r.font.color.rgb = c_dark_slate
    else:
        p_empty = tf_e1.add_paragraph()
        p_empty.text = "• No projects are currently in Red or Amber status requiring immediate escalations."
        p_empty.font.name = "Arial"
        p_empty.font.size = Pt(10.5)
        p_empty.font.color.rgb = c_medium_slate

    # Right Column Card: General Portfolio Recommendations
    card_e2 = add_card(slide_actions, Inches(6.88), Inches(1.35), Inches(5.9), Inches(5.2))
    tf_e2 = card_e2.text_frame
    format_tf(tf_e2)

    p_e2_hdr = tf_e2.paragraphs[0]
    p_e2_hdr.text = "GENERAL PORTFOLIO GOVERNANCE"
    p_e2_hdr.font.name = "Arial"
    p_e2_hdr.font.size = Pt(12)
    p_e2_hdr.font.bold = True
    p_e2_hdr.font.color.rgb = c_accent_indigo
    p_e2_hdr.space_after = Pt(10)

    general_recs = [
        "Enforce Weekly Baseline Syncs: Ensure all project managers update schedule baselines weekly to maintain high-quality tracking metrics.",
        "Address Critical Negative Float: Conduct dedicated review sessions for tasks exhibiting negative float to resolve blockers before milestone slips occur.",
        "Refine Commenting Standard: Ensure PM comment updates describe concrete mitigation steps rather than status summaries to keep sentiment scores clear.",
        "Verify Workbook Data Quality: Fix missing ownership and baseline dates inside source spreadsheets to increase data quality confidence scores.",
        "Establish Trend Milestones: Leverage SQLite time-series data to compare current month trends against baseline trajectory projections."
    ]

    for rec in general_recs:
        p_rec = tf_e2.add_paragraph()
        p_rec.space_after = Pt(8)
        run_bullet = p_rec.add_run()
        run_bullet.text = "• "
        run_bullet.font.name = "Arial"
        run_bullet.font.bold = True
        run_bullet.font.size = Pt(10)
        run_bullet.font.color.rgb = c_accent_indigo

        run_text = p_rec.add_run()
        parts = rec.split(": ", 1)
        if len(parts) == 2:
            run_title = p_rec.add_run()
            run_title.text = parts[0] + ": "
            run_title.font.name = "Arial"
            run_title.font.bold = True
            run_title.font.size = Pt(10)
            run_title.font.color.rgb = c_dark_slate
            
            run_desc = p_rec.add_run()
            run_desc.text = parts[1]
            run_desc.font.name = "Arial"
            run_desc.font.size = Pt(10)
            run_desc.font.color.rgb = c_medium_slate
        else:
            run_text.text = rec
            run_text.font.name = "Arial"
            run_text.font.size = Pt(10)
            run_text.font.color.rgb = c_dark_slate

    # ----------------------------------------------------
    # Slide 6: How Project Works & Solving the Problem
    # ----------------------------------------------------
    slide_method = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide_method, "How Project Works: RAG Engine & Problem-Solving", "Standard RAG signal weights, automatic normalization, and decision auditability", "How Project Works")

    # Left Column Card: Standard Signal Weights & Logic Model
    card_m1 = add_card(slide_method, Inches(0.55), Inches(1.35), Inches(5.9), Inches(5.2))
    tf_m1 = card_m1.text_frame
    format_tf(tf_m1)

    p_m1_hdr = tf_m1.paragraphs[0]
    p_m1_hdr.text = "DETERMINISTIC RAG LOGIC MODEL"
    p_m1_hdr.font.name = "Arial"
    p_m1_hdr.font.size = Pt(12)
    p_m1_hdr.font.bold = True
    p_m1_hdr.font.color.rgb = c_accent_indigo
    p_m1_hdr.space_after = Pt(8)

    model_points = [
        "Schedule Health (30%): Evaluates task delays, late critical paths, float erosion.",
        "Progress Gap (20%): Compares actual % complete vs expected progress by date.",
        "Milestones Confidence (20%): Measures completed vs delayed key milestones.",
        "Blocker Keywords (15%): Scans task names and status comments for blocker keywords.",
        "Stakeholder Sentiment (10%): NLP-analyzed tone and outlook of weekly PM comments.",
        "Budget Variance (5%): Burn rates and cost variance (omitted if data is absent)."
    ]
    for pt in model_points:
        p_pt = tf_m1.add_paragraph()
        p_pt.space_after = Pt(6)
        
        parts = pt.split(": ", 1)
        run_bullet = p_pt.add_run()
        run_bullet.text = "• "
        run_bullet.font.name = "Arial"
        run_bullet.font.bold = True
        run_bullet.font.size = Pt(9.5)
        run_bullet.font.color.rgb = c_accent_indigo

        run_title = p_pt.add_run()
        run_title.text = parts[0] + ": "
        run_title.font.name = "Arial"
        run_title.font.bold = True
        run_title.font.size = Pt(9.5)
        run_title.font.color.rgb = c_dark_slate

        run_desc = p_pt.add_run()
        run_desc.text = parts[1]
        run_desc.font.name = "Arial"
        run_desc.font.size = Pt(9.5)
        run_desc.font.color.rgb = c_medium_slate

    p_over = tf_m1.add_paragraph()
    p_over.space_before = Pt(8)
    p_over.space_after = Pt(4)
    run_over_title = p_over.add_run()
    run_over_title.text = "Critical Override Rules:"
    run_over_title.font.name = "Arial"
    run_over_title.font.bold = True
    run_over_title.font.size = Pt(10.5)
    run_over_title.font.color.rgb = c_red

    p_over_body = tf_m1.add_paragraph()
    p_over_body.text = "If any critical task has a Red schedule health and is either root-level or at-risk level is High, the overall project status is forced to Red regardless of other signals. Ensures major threats are never hidden."
    p_over_body.font.name = "Arial"
    p_over_body.font.size = Pt(9)
    p_over_body.font.color.rgb = c_medium_slate

    # Right Column Card: Efficiency & Solving the Problem
    card_m2 = add_card(slide_method, Inches(6.88), Inches(1.35), Inches(5.9), Inches(5.2))
    tf_m2 = card_m2.text_frame
    format_tf(tf_m2)

    p_m2_hdr = tf_m2.paragraphs[0]
    p_m2_hdr.text = "EFFICIENCY & PROBLEM-SOLVING"
    p_m2_hdr.font.name = "Arial"
    p_m2_hdr.font.size = Pt(12)
    p_m2_hdr.font.bold = True
    p_m2_hdr.font.color.rgb = c_accent_indigo
    p_m2_hdr.space_after = Pt(8)

    efficiency_points = [
        "Eliminating Reporting Bias: Replaces subjective/manual project health reporting with automated, uniform, data-driven evaluation.",
        "Automatic Data Normalization: Cleans and parses diverse Excel formats. Automatically infers active project stages from task hierarchies and dates.",
        "Uncompromised Auditability: Every AI narrative and signal cites precise Excel row numbers, allowing stakeholders to trace results directly back to the sheet.",
        "Resilient Fallbacks: Missing baseline dates default to current dates; missing comments default to a neutral sentiment score, ensuring the pipeline never fails.",
        "Historical Persistence: Stores all weekly runs in a unified SQLite database, enabling trend analytics and automated RAG flip alerts."
    ]

    for pt in efficiency_points:
        p_pt = tf_m2.add_paragraph()
        p_pt.space_after = Pt(6)
        
        parts = pt.split(": ", 1)
        run_bullet = p_pt.add_run()
        run_bullet.text = "• "
        run_bullet.font.name = "Arial"
        run_bullet.font.bold = True
        run_bullet.font.size = Pt(9.5)
        run_bullet.font.color.rgb = c_accent_indigo

        run_title = p_pt.add_run()
        run_title.text = parts[0] + ": "
        run_title.font.name = "Arial"
        run_title.font.bold = True
        run_title.font.size = Pt(9.5)
        run_title.font.color.rgb = c_dark_slate

        run_desc = p_pt.add_run()
        run_desc.text = parts[1]
        run_desc.font.name = "Arial"
        run_desc.font.size = Pt(9.5)
        run_desc.font.color.rgb = c_medium_slate

    # Add Footer & Slide Number on every slide
    for idx, slide in enumerate(prs.slides, start=1):
        footer = slide.shapes.add_textbox(Inches(0.55), Inches(7.05), Inches(6.0), Inches(0.3))
        p_f = footer.text_frame.paragraphs[0]
        p_f.text = "Zycus Project Health reporting agent | Confidential"
        p_f.font.name = "Arial"
        p_f.font.size = Pt(8.5)
        p_f.font.color.rgb = c_medium_slate

        num_box = slide.shapes.add_textbox(Inches(11.78), Inches(7.05), Inches(1.0), Inches(0.3))
        p_n = num_box.text_frame.paragraphs[0]
        p_n.text = f"Slide {idx} of {len(prs.slides)}"
        p_n.alignment = PP_ALIGN.RIGHT
        p_n.font.name = "Arial"
        p_n.font.size = Pt(8.5)
        p_n.font.color.rgb = c_medium_slate

    prs.save(pptx_path)


def _theme_bullets(comments: list[str]) -> list[str]:
    if not comments:
        return ["No comment themes available from stored snapshots."]
    themes = {
        "Data/mapping dependency": ["mapping", "data", "sample", "jde"],
        "Workshop or schedule impact": ["workshop", "scheduled", "dates", "impacted"],
        "Open owner follow-up": ["need", "@", "pending", "remain"],
    }
    bullets = []
    text = "\n".join(comments).lower()
    for theme, words in themes.items():
        count = sum(text.count(word) for word in words)
        if count:
            bullets.append(f"{theme}: {count} keyword hits")
    return bullets or ["Comments are present but do not cluster around known risk themes."]


def _deck_markdown(conn: sqlite3.Connection) -> str:
    portfolio = fetch_portfolio(conn)
    risks = fetch_top_risk_tasks(conn, 8)
    lines = ["# Project Health Executive Review", "", "## Portfolio Health Overview"]
    lines.extend([f"- {row['name']}: {row['rag_status']} ({row['rag_score']}/100)" for row in portfolio])
    lines.extend(["", "## Emerging Risks"])
    lines.extend([f"- {row['name']}: row {row['source_row']} - {row['task_name']}" for row in risks])
    lines.extend(["", "## Recommendations", "- Use `pip install -r requirements.txt` to enable PPTX generation with python-pptx."])
    return "\n".join(lines)
